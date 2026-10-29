#!/usr/bin/env python3
"""Automated screenshot capture suite for ConfMan project report Appendix D.

Captures:
- Figure D.1: Dashboard - Home Page & Login Screen
- Figure D.2: Chair Dashboard - KPI Cards & Submission Queue
- Figure D.3: Paper Detail View with Review Panel & Suggestion
- Figure D.4: Payment Checkout & Registration Confirmation

Embeds the generated screenshots directly into the DOCX reports.
"""
import os
import sys
import threading
import time
from PIL import Image
import docx
from docx.shared import Inches

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

from cms_app import create_app
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By


def combine_images_vertical(img1_path, img2_path, out_path, spacing=15):
    im1 = Image.open(img1_path)
    im2 = Image.open(img2_path)
    w = max(im1.width, im2.width)
    h = im1.height + im2.height + spacing
    combined = Image.new("RGB", (w, h), (245, 247, 248))
    combined.paste(im1, (0, 0))
    combined.paste(im2, (0, im1.height + spacing))
    combined.save(out_path, quality=95)
    return out_path


def main():
    out_dir = os.path.join(PROJECT_ROOT, "docs", "screenshots")
    os.makedirs(out_dir, exist_ok=True)

    print("1. Starting local ConfMan Flask instance...")
    app = create_app({"AUTO_SEED": True, "CMS_SECRET": "report-screenshots-secret"})
    server = threading.Thread(
        target=lambda: app.run(host="127.0.0.1", port=5065, use_reloader=False),
        daemon=True,
    )
    server.start()
    time.sleep(2)

    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1440,950")

    driver = webdriver.Chrome(options=options)
    base_url = "http://127.0.0.1:5065"

    def login_as(email):
        driver.delete_all_cookies()
        driver.get(f"{base_url}/login")
        time.sleep(1)
        btn = driver.find_element(By.CSS_SELECTOR, f'button[data-email="{email}"]')
        btn.click()
        time.sleep(1.5)

    # --- Figure D.1: Home Page & Login Screen ---
    print("2. Capturing Figure D.1 (Home & Login Screen)...")
    driver.delete_all_cookies()
    driver.get(f"{base_url}/")
    time.sleep(1)
    home_img = os.path.join(out_dir, "figure_d1_home.png")
    driver.save_screenshot(home_img)

    driver.get(f"{base_url}/login")
    time.sleep(1)
    login_img = os.path.join(out_dir, "figure_d1_login.png")
    driver.save_screenshot(login_img)

    fig_d1 = os.path.join(out_dir, "Figure_D1_Dashboard_Home_and_Login.png")
    combine_images_vertical(home_img, login_img, fig_d1)
    print(f"  ✓ Saved: {fig_d1}")

    # --- Figure D.2: Chair Dashboard ---
    print("3. Capturing Figure D.2 (Chair Dashboard)...")
    login_as("chair@demo.org")
    driver.get(f"{base_url}/dashboard")
    time.sleep(1.5)
    fig_d2 = os.path.join(out_dir, "Figure_D2_Chair_Dashboard.png")
    driver.save_screenshot(fig_d2)
    print(f"  ✓ Saved: {fig_d2}")

    # --- Figure D.3: Paper Detail View with Review Panel ---
    print("4. Capturing Figure D.3 (Paper Detail View with Review Panel)...")
    driver.get(f"{base_url}/papers/10")
    time.sleep(1.5)
    driver.execute_script("window.scrollTo(0, 320);")
    time.sleep(0.5)
    fig_d3 = os.path.join(out_dir, "Figure_D3_Paper_Detail_Review_Panel.png")
    driver.save_screenshot(fig_d3)
    print(f"  ✓ Saved: {fig_d3}")

    # --- Figure D.4: Payment Checkout & Registration Confirmation ---
    print("5. Capturing Figure D.4 (Payment Checkout & Confirmation)...")
    login_as("attendee@demo.org")

    # Payment checkout page for pending registration (reg_id=7)
    driver.get(f"{base_url}/registrations/7/pay")
    time.sleep(1)
    pay_img = os.path.join(out_dir, "figure_d4_checkout.png")
    driver.save_screenshot(pay_img)

    # Submit mock card
    card_input = driver.find_element(By.NAME, "card")
    card_input.clear()
    card_input.send_keys("4242 4242 4242 4242")
    driver.find_element(By.CSS_SELECTOR, "button.btn.amber").click()
    time.sleep(1.5)

    # Attendee dashboard showing confirmed registration
    confirm_img = os.path.join(out_dir, "figure_d4_confirmed.png")
    driver.save_screenshot(confirm_img)

    fig_d4 = os.path.join(out_dir, "Figure_D4_Payment_Checkout_and_Confirmation.png")
    combine_images_vertical(pay_img, confirm_img, fig_d4)
    print(f"  ✓ Saved: {fig_d4}")

    driver.quit()

    # --- Embed into Word Documents ---
    print("\n6. Embedding screenshots into reports...")
    report_paths = [
        os.path.join(PROJECT_ROOT, "docs", "reports", "confman_report_refined.docx"),
        os.path.join(PROJECT_ROOT, "docs", "reports", "confman_report.docx"),
    ]

    figures = {
        "[PLACEHOLDER – Figure D.1]": fig_d1,
        "[PLACEHOLDER – Figure D.2]": fig_d2,
        "[PLACEHOLDER – Figure D.3]": fig_d3,
        "[PLACEHOLDER – Figure D.4]": fig_d4,
    }

    for report_file in report_paths:
        if not os.path.exists(report_file):
            continue
        print(f"  Updating: {os.path.basename(report_file)}")
        doc = docx.Document(report_file)
        modified = False

        for i, p in enumerate(doc.paragraphs):
            for placeholder, img_path in figures.items():
                if placeholder in p.text:
                    p.text = p.text.replace(placeholder, "").strip()
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(6.0))
                    modified = True
                    print(f"    ✓ Embedded {os.path.basename(img_path)} into {os.path.basename(report_file)}")

        if modified:
            doc.save(report_file)
            print(f"    ✓ Successfully saved {os.path.basename(report_file)}")

    print("\nAll screenshot tasks completed successfully!")


if __name__ == "__main__":
    main()
